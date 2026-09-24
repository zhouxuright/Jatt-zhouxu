"""Generate test files (TXT/MD/PDF/DOCX) for batch upload E2E testing."""
import os

OUT = os.path.join(os.path.dirname(__file__), "test_files")
os.makedirs(OUT, exist_ok=True)

TITLE = "劳动合同解除与经济补偿金法律要点"
BODY = """劳动合同解除与经济补偿金的法律要点说明。

根据《中华人民共和国劳动合同法》第四十六条规定，有下列情形之一的，用人单位应当向劳动者支付经济补偿：
（一）劳动者依照本法第三十八条规定解除劳动合同的；
（二）用人单位依照本法第三十六条规定向劳动者提出解除劳动合同并与劳动者协商一致解除劳动合同的；
（三）用人单位依照本法第四十条规定解除劳动合同的；
（四）用人单位依照本法第四十一条第一款规定解除劳动合同的。

根据第四十七条规定，经济补偿按劳动者在本单位工作的年限，每满一年支付一个月工资的标准向劳动者支付。六个月以上不满一年的，按一年计算；不满六个月的，向劳动者支付半个月工资的经济补偿。

劳动者月工资高于用人单位所在直辖市、设区的市级人民政府公布的本地区上年度职工月平均工资三倍的，向其支付经济补偿的标准按职工月平均工资三倍的数额支付，向其支付经济补偿的年限最高不超过十二年。

用人单位违法解除或者终止劳动合同，劳动者要求继续履行劳动合同的，用人单位应当继续履行；劳动者不要求继续履行劳动合同或者劳动合同已经不能继续履行的，用人单位应当依照本法第八十七条规定支付赔偿金，即经济补偿标准的二倍。
"""

MARKDOWN = f"""# 劳动合同解除与经济补偿金法律要点

## 一、支付经济补偿的情形

根据《**中华人民共和国劳动合同法**》第四十六条，以下情形用人单位应当支付经济补偿：

- 劳动者依法解除（第三十八条）
- 用人单位提出并协商一致解除（第三十六条）
- 无过失性辞退（第四十条）
- 经济性裁员（第四十一条第一款）

## 二、经济补偿的计算标准

> 按劳动者在本单位工作的年限，**每满一年支付一个月工资**；六个月以上不满一年按一年计算；不满六个月支付半个月工资。（第四十七条）

## 三、高收入劳动者的上限

月工资高于当地上年度职工月平均工资三倍的，按三倍封顶，年限最高十二年。

## 四、违法解除的赔偿金

违法解除或终止劳动合同的，按第八十七条支付**赔偿金（经济补偿标准二倍）**。
"""


def write_txt(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write(TITLE + "\n\n" + BODY)


def write_md(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write(MARKDOWN)


def write_pdf(path):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("STSong-Light", 16)
    c.drawString(72, 800, TITLE)
    c.setFont("STSong-Light", 11)
    y = 770
    for line in BODY.replace("\n\n", "\n").split("\n"):
        if y < 60:
            c.showPage()
            c.setFont("STSong-Light", 11)
            y = 800
        # wrap long lines
        while len(line) > 45:
            c.drawString(72, y, line[:45])
            line = line[45:]
            y -= 16
        c.drawString(72, y, line)
        y -= 16
    c.save()


def write_docx(path):
    import docx
    d = docx.Document()
    d.add_heading(TITLE, level=1)
    for para in BODY.split("\n\n"):
        d.add_paragraph(para)
    d.save(path)


if __name__ == "__main__":
    write_txt(os.path.join(OUT, "test_law_contract.txt"))
    write_md(os.path.join(OUT, "test_law_contract.md"))
    write_pdf(os.path.join(OUT, "test_law_contract.pdf"))
    write_docx(os.path.join(OUT, "test_law_contract.docx"))
    print("Generated test files in", OUT)
    for f in sorted(os.listdir(OUT)):
        p = os.path.join(OUT, f)
        print(f"  {f}: {os.path.getsize(p)} bytes")