"""
深度文档解析服务 — 法律文档结构化提取

支持格式: PDF / DOCX / DOC / TXT / Markdown
提取内容: 文本 + 结构(章节/标题层级) + 表格 + 元数据 + 法律要素

优雅降级:
- 无 PyMuPDF → 使用 pdfminer 或纯文本回退
- 无 python-docx → 使用 zipfile + XML 基础解析
- 所有解析器均可在最小依赖下工作

使用方法:
    from app.services.document_parser import DocumentParser
    parser = DocumentParser()
    result = await parser.parse("/path/to/contract.pdf", "pdf")
"""

from __future__ import annotations

import logging
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)


# ============================================================================
# Data Models
# ============================================================================

@dataclass
class Section:
    """文档章节"""
    level: int = 0
    title: str = ""
    content: str = ""
    page: int = 0
    children: list["Section"] = field(default_factory=list)


@dataclass
class Table:
    """表格"""
    page: int = 0
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    caption: str = ""


@dataclass
class LegalElements:
    """从文档中提取的法律要素"""
    dates: list[str] = field(default_factory=list)          # 日期
    parties: list[str] = field(default_factory=list)         # 当事人
    amounts: list[str] = field(default_factory=list)         # 金额
    case_numbers: list[str] = field(default_factory=list)    # 案号
    law_references: list[str] = field(default_factory=list)  # 法条引用
    addresses: list[str] = field(default_factory=list)       # 地址
    organizations: list[str] = field(default_factory=list)   # 机构名称


@dataclass
class ParsedDocument:
    """解析后的文档"""
    text: str = ""
    sections: list[Section] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    legal_elements: LegalElements = field(default_factory=LegalElements)
    parse_success: bool = True
    parse_errors: list[str] = field(default_factory=list)
    file_type: str = ""
    char_count: int = 0
    page_count: int = 0


# ============================================================================
# Document Parser
# ============================================================================

class DocumentParser:
    """深度文档解析器"""

    async def parse(self, file_path: str, file_type: str | None = None) -> ParsedDocument:
        """
        解析文档文件，返回结构化结果。

        Args:
            file_path: 文件路径
            file_type: 文件类型 (pdf/docx/doc/txt/md)。None则自动检测。

        Returns:
            ParsedDocument: 解析结果
        """
        path = Path(file_path)
        if not path.exists():
            return ParsedDocument(
                parse_success=False,
                parse_errors=[f"文件不存在: {file_path}"],
            )

        if file_type is None:
            file_type = self._detect_type(path)

        logger.info("Parsing %s as %s (%.1f KB)", path.name, file_type, path.stat().st_size / 1024)

        try:
            if file_type == "pdf":
                return self._parse_pdf(path)
            elif file_type == "docx":
                return self._parse_docx(path)
            elif file_type == "doc":
                return self._parse_doc(path)
            elif file_type == "md" or file_type == "markdown":
                return self._parse_markdown(path)
            else:
                return self._parse_txt(path)
        except Exception as e:
            logger.exception("Failed to parse %s: %s", file_path, e)
            return ParsedDocument(
                parse_success=False,
                parse_errors=[str(e)],
                file_type=file_type,
            )

    def _detect_type(self, path: Path) -> str:
        suffix = path.suffix.lower()
        mapping = {
            ".pdf": "pdf",
            ".docx": "docx",
            ".doc": "doc",
            ".txt": "txt",
            ".md": "md",
            ".markdown": "md",
        }
        return mapping.get(suffix, "txt")

    # ------------------------------------------------------------------
    # PDF Parser
    # ------------------------------------------------------------------

    def _parse_pdf(self, path: Path) -> ParsedDocument:
        """使用 PyMuPDF 解析 PDF"""
        try:
            import fitz  # PyMuPDF
            return self._parse_pdf_with_fitz(path)
        except ImportError:
            logger.warning("PyMuPDF not installed, using basic PDF parsing")
            return self._parse_pdf_basic(path)

    def _parse_pdf_with_fitz(self, path: Path) -> ParsedDocument:
        """PyMuPDF 深度解析"""
        import fitz

        doc = fitz.open(str(path))
        pages_text = []
        sections = []
        tables = []

        for page_num in range(len(doc)):
            page = doc[page_num]

            # 提取文本 (保留布局)
            text = page.get_text("text")
            pages_text.append(text)

            # 提取表格
            try:
                tab_finder = page.find_tables()
                for tab in tab_finder.tables:
                    table_data = tab.extract()
                    if table_data and len(table_data) > 0:
                        headers = [str(c) if c else "" for c in table_data[0]]
                        rows = [[str(c) if c else "" for c in row] for row in table_data[1:]]
                        tables.append(Table(
                            page=page_num + 1,
                            headers=headers,
                            rows=rows,
                        ))
            except Exception as e:
                logger.debug("Table extraction failed on page %d: %s", page_num, e)

            # 提取章节标题 (基于字体大小)
            try:
                blocks = page.get_text("dict")["blocks"]
                for block in blocks:
                    if block.get("type") == 0:  # text block
                        for line in block.get("lines", []):
                            for span in line.get("spans", []):
                                font_size = span.get("size", 0)
                                txt = span.get("text", "").strip()
                                if font_size > 14 and txt and len(txt) < 100:
                                    level = 1 if font_size > 18 else 2
                                    sections.append(Section(
                                        level=level,
                                        title=txt,
                                        page=page_num + 1,
                                    ))
            except Exception:
                pass

        full_text = "\n".join(pages_text)
        metadata = {
            "page_count": len(doc),
            "title": doc.metadata.get("title", ""),
            "author": doc.metadata.get("author", ""),
            "creator": doc.metadata.get("creator", ""),
            "creation_date": str(doc.metadata.get("creationDate", "")),
            "modification_date": str(doc.metadata.get("modDate", "")),
        }
        doc.close()

        legal_elements = self._extract_legal_elements(full_text)

        return ParsedDocument(
            text=full_text,
            sections=sections,
            tables=tables,
            metadata=metadata,
            legal_elements=legal_elements,
            file_type="pdf",
            char_count=len(full_text),
            page_count=len(pages_text),
        )

    def _parse_pdf_basic(self, path: Path) -> ParsedDocument:
        """基础PDF解析 (无PyMuPDF依赖)"""
        try:
            # 尝试 pdfminer
            from pdfminer.high_level import extract_text
            text = extract_text(str(path))
            legal_elements = self._extract_legal_elements(text)
            return ParsedDocument(
                text=text,
                metadata={"parser": "pdfminer"},
                legal_elements=legal_elements,
                file_type="pdf",
                char_count=len(text),
            )
        except ImportError:
            pass

        # 最后回退: 尝试直接读取 (对某些简单PDF有效)
        try:
            with open(path, "rb") as f:
                content = f.read()
            # 提取可打印ASCII文本
            text = "".join(chr(b) for b in content if 32 <= b < 127 or b in (10, 13))
            legal_elements = self._extract_legal_elements(text)
            return ParsedDocument(
                text=text,
                metadata={"parser": "raw_binary"},
                legal_elements=legal_elements,
                file_type="pdf",
                char_count=len(text),
                parse_errors=["PyMuPDF和pdfminer均未安装，使用原始二进制提取"],
            )
        except Exception as e:
            return ParsedDocument(
                parse_success=False,
                parse_errors=[f"PDF解析失败: {e}"],
                file_type="pdf",
            )

    # ------------------------------------------------------------------
    # DOCX Parser
    # ------------------------------------------------------------------

    def _parse_docx(self, path: Path) -> ParsedDocument:
        """解析 DOCX 文件"""
        try:
            from docx import Document
            return self._parse_docx_with_python_docx(path)
        except ImportError:
            logger.warning("python-docx not installed, using basic DOCX parsing")
            return self._parse_docx_basic(path)

    def _parse_docx_with_python_docx(self, path: Path) -> ParsedDocument:
        """使用 python-docx 深度解析"""
        from docx import Document

        doc = Document(str(path))
        paragraphs = []
        sections = []
        tables = []

        # 提取段落
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
            paragraphs.append(text)

            # 检测标题
            if para.style and para.style.name:
                style_name = para.style.name.lower()
                if "heading" in style_name:
                    try:
                        level = int(para.style.name.replace("Heading", "").strip())
                    except ValueError:
                        level = 1
                    sections.append(Section(level=level, title=text))

        # 提取表格
        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            if rows:
                tables.append(Table(
                    headers=rows[0] if rows else [],
                    rows=rows[1:] if len(rows) > 1 else [],
                ))

        full_text = "\n".join(paragraphs)

        # 提取元数据
        metadata = {}
        if doc.core_properties:
            props = doc.core_properties
            metadata = {
                "title": props.title or "",
                "author": props.author or "",
                "created": str(props.created) if props.created else "",
                "modified": str(props.modified) if props.modified else "",
            }

        legal_elements = self._extract_legal_elements(full_text)

        return ParsedDocument(
            text=full_text,
            sections=sections,
            tables=tables,
            metadata=metadata,
            legal_elements=legal_elements,
            file_type="docx",
            char_count=len(full_text),
            page_count=len(doc.sections),
        )

    def _parse_docx_basic(self, path: Path) -> ParsedDocument:
        """基础DOCX解析 (zipfile + XML)"""
        import zipfile
        import xml.etree.ElementTree as ET

        try:
            with zipfile.ZipFile(path) as z:
                with z.open("word/document.xml") as f:
                    tree = ET.parse(f)

            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = []
            for p in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                texts = [t.text for t in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t") if t.text]
                if texts:
                    paragraphs.append("".join(texts))

            full_text = "\n".join(paragraphs)
            legal_elements = self._extract_legal_elements(full_text)

            return ParsedDocument(
                text=full_text,
                metadata={"parser": "zipfile_xml"},
                legal_elements=legal_elements,
                file_type="docx",
                char_count=len(full_text),
            )
        except Exception as e:
            return ParsedDocument(
                parse_success=False,
                parse_errors=[f"DOCX解析失败: {e}"],
                file_type="docx",
            )

    # ------------------------------------------------------------------
    # DOC Parser (legacy format)
    # ------------------------------------------------------------------

    def _parse_doc(self, path: Path) -> ParsedDocument:
        """解析旧版 DOC 文件 — 尝试多种方式"""
        # 方法1: 直接文本提取 (很多.doc文件实际上包含可提取文本)
        try:
            with open(path, "rb") as f:
                content = f.read()

            # 尝试提取 UTF-8 和 GBK 文本
            text_parts = []
            for encoding in ["utf-8", "gbk", "gb2312", "gb18030"]:
                try:
                    decoded = content.decode(encoding, errors="ignore")
                    # 过滤控制字符
                    cleaned = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', decoded)
                    if len(cleaned) > 100:
                        text_parts.append((len(cleaned), cleaned))
                except Exception:
                    continue

            if text_parts:
                text_parts.sort(reverse=True)
                full_text = text_parts[0][1]
                legal_elements = self._extract_legal_elements(full_text)
                return ParsedDocument(
                    text=full_text[:500000],  # 限制大小
                    metadata={"parser": "binary_text_extraction"},
                    legal_elements=legal_elements,
                    file_type="doc",
                    char_count=len(full_text),
                    parse_errors=["DOC文件使用基础文本提取，建议转换为DOCX格式以获得更好效果"],
                )
        except Exception as e:
            pass

        return ParsedDocument(
            parse_success=False,
            parse_errors=["DOC文件解析失败，请转换为DOCX或PDF格式"],
            file_type="doc",
        )

    # ------------------------------------------------------------------
    # TXT / Markdown Parser
    # ------------------------------------------------------------------

    def _parse_txt(self, path: Path) -> ParsedDocument:
        """解析纯文本文件"""
        # 自动检测编码
        for encoding in ["utf-8", "gbk", "gb2312", "gb18030", "latin-1"]:
            try:
                with open(path, "r", encoding=encoding) as f:
                    text = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue
        else:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()

        # 检测章节结构
        sections = self._detect_sections_from_text(text)
        legal_elements = self._extract_legal_elements(text)

        return ParsedDocument(
            text=text,
            sections=sections,
            metadata={"parser": "text_auto_encoding"},
            legal_elements=legal_elements,
            file_type="txt",
            char_count=len(text),
        )

    def _parse_markdown(self, path: Path) -> ParsedDocument:
        """解析 Markdown 文件"""
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()

        # 提取 Markdown 标题作为章节
        sections = []
        for match in re.finditer(r'^(#{1,6})\s+(.+)$', text, re.MULTILINE):
            level = len(match.group(1))
            title = match.group(2).strip()
            sections.append(Section(level=level, title=title))

        # 提取表格 (Markdown 格式)
        tables = []
        table_pattern = re.finditer(
            r'\|(.+)\|\n\|[-| :]+\|\n((?:\|.+\|\n?)+)',
            text
        )
        for match in table_pattern:
            headers = [h.strip() for h in match.group(1).split('|') if h.strip()]
            rows = []
            for row_line in match.group(2).strip().split('\n'):
                cells = [c.strip() for c in row_line.split('|') if c.strip()]
                rows.append(cells)
            tables.append(Table(headers=headers, rows=rows))

        legal_elements = self._extract_legal_elements(text)

        return ParsedDocument(
            text=text,
            sections=sections,
            tables=tables,
            metadata={"parser": "markdown"},
            legal_elements=legal_elements,
            file_type="md",
            char_count=len(text),
        )

    # ------------------------------------------------------------------
    # Legal Element Extraction
    # ------------------------------------------------------------------

    def _extract_legal_elements(self, text: str) -> LegalElements:
        """从文本中提取法律要素"""
        elements = LegalElements()

        if not text:
            return elements

        # 日期提取
        date_patterns = [
            r'\d{4}年\d{1,2}月\d{1,2}日',
            r'\d{4}[-/]\d{1,2}[-/]\d{1,2}',
            r'\d{4}年',
        ]
        for pattern in date_patterns:
            elements.dates.extend(re.findall(pattern, text))
        elements.dates = list(set(elements.dates))[:50]

        # 案号提取
        case_no_patterns = [
            r'[（(]\d{4}[)）][^\s]{2,20}第\s*\d+\s*号',
            r'[（(]\d{4}[)）][^\s]{2,20}\d+号',
        ]
        for pattern in case_no_patterns:
            elements.case_numbers.extend(re.findall(pattern, text))
        elements.case_numbers = list(set(elements.case_numbers))[:20]

        # 法条引用
        law_ref_patterns = [
            r'《[^》]+》第[一二三四五六七八九十百千零\d]+条[第款各项]?',
            r'《[^》]+》',
        ]
        for pattern in law_ref_patterns:
            elements.law_references.extend(re.findall(pattern, text))
        elements.law_references = list(set(elements.law_references))[:50]

        # 金额提取
        amount_patterns = [
            r'人民币\s*[\d,]+(?:\.\d+)?\s*[亿万]?元',
            r'[\d,]+(?:\.\d+)?\s*[亿万]元',
            r'[\d,]+(?:\.\d+)?\s*元',
        ]
        for pattern in amount_patterns:
            elements.amounts.extend(re.findall(pattern, text))
        elements.amounts = list(set(elements.amounts))[:30]

        # 当事人 (简化提取 — 原告/被告/申请人/被申请人)
        party_patterns = [
            r'(?:原告|申请人)[：:]\s*([^\s，,。]+)',
            r'(?:被告|被申请人)[：:]\s*([^\s，,。]+)',
        ]
        for pattern in party_patterns:
            elements.parties.extend(re.findall(pattern, text))
        elements.parties = list(set(elements.parties))[:20]

        return elements

    def _detect_sections_from_text(self, text: str) -> list[Section]:
        """从纯文本中检测章节结构"""
        sections = []
        # 中文数字编号: 第一章、第二节、第三条
        patterns = [
            (1, r'^第[一二三四五六七八九十百]+[章节编篇]\s+.*$'),
            (2, r'^第[一二三四五六七八九十百]+条\s*$'),
            (2, r'^\d+\.\d+\s+.*$'),
            (3, r'^\d+\.\d+\.\d+\s+.*$'),
        ]
        for level, pattern in patterns:
            for match in re.finditer(pattern, text, re.MULTILINE):
                sections.append(Section(level=level, title=match.group().strip()))

        return sections


# ============================================================================
# Singleton
# ============================================================================

_parser_instance: DocumentParser | None = None


def get_document_parser() -> DocumentParser:
    """获取文档解析器单例"""
    global _parser_instance
    if _parser_instance is None:
        _parser_instance = DocumentParser()
    return _parser_instance
